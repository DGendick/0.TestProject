import pandas as pd
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from collections import defaultdict
import re

class DocumentRemarksCounter:
    def __init__(self):
        self.remarks_data = defaultdict(lambda: {
            'total': 0,
            'by_status': defaultdict(int),
            'details': []
        })
        self.total_remarks = 0
        self.total_files = 0
        self.file_list = []
    
    def classify_remark_by_content(self, comment_text):
        """Классификация замечаний по содержанию"""
        text_lower = comment_text.lower()
        
        if any(word in text_lower for word in ['откорректировать', 'исправить', 'изменить']):
            return 'Требуется корректировка'
        elif any(word in text_lower for word in ['дублирует', 'повтор']):
            return 'Дублирование информации'
        elif any(word in text_lower for word in ['пронумеровать', 'нумерация', 'исключить']):
            return 'Оформление и нумерация'
        elif any(word in text_lower for word in ['подключение', 'схема', 'кабель']):
            return 'Ошибки в схемах подключения'
        elif any(word in text_lower for word in ['наименование', 'идентификатор', 'аналогии']):
            return 'Некорректное наименование'
        else:
            return 'Прочие замечания'
    
    def parse_comment_review_sheet(self, file_path):
        """Парсинг файла формата Comment Review Sheet"""
        try:
            df = pd.read_excel(file_path, sheet_name='Comment Review Sheet', header=None)
            
            # Находим строку с заголовками
            header_row = None
            for idx, row in df.iterrows():
                if row.astype(str).str.contains('Comment No', case=False, na=False).any():
                    header_row = idx
                    break
            
            if header_row is None:
                return []
            
            headers = df.iloc[header_row].astype(str).str.strip()
            col_index = {}
            
            for idx, header in headers.items():
                header_lower = header.lower()
                if 'comment no' in header_lower:
                    col_index['comment_no'] = idx
                elif 'owner comments in russian' in header_lower:
                    col_index['comment'] = idx
                elif 'incorporation status' in header_lower:
                    col_index['status'] = idx
                elif "contractor's acceptance" in header_lower:
                    col_index['acceptance'] = idx
            
            if 'comment' not in col_index:
                return []
            
            data_start = header_row + 1
            file_remarks = []
            
            for idx in range(data_start, len(df)):
                row = df.iloc[idx]
                
                comment_text = str(row[col_index['comment']]) if col_index['comment'] < len(row) else ''
                if pd.isna(comment_text) or comment_text == 'nan' or not comment_text.strip():
                    continue
                
                comment_no = str(row[col_index['comment_no']]) if 'comment_no' in col_index and col_index['comment_no'] < len(row) else ''
                status = str(row[col_index['status']]) if 'status' in col_index and col_index['status'] < len(row) else ''
                status = 'Open' if 'open' in status.lower() else 'Closed' if 'closed' in status.lower() else 'Unknown'
                acceptance = str(row[col_index['acceptance']]) if 'acceptance' in col_index and col_index['acceptance'] < len(row) else ''
                
                remark_type = self.classify_remark_by_content(comment_text)
                
                file_remarks.append({
                    'comment_no': comment_no,
                    'text': comment_text.strip(),
                    'type': remark_type,
                    'status': status,
                    'acceptance': acceptance
                })
            
            return file_remarks
            
        except Exception as e:
            print(f"  Ошибка при обработке файла {file_path.name}: {e}")
            return []
    
    def process_folder(self, folder_path):
        """Обработка всех файлов в папке"""
        excel_files = list(Path(folder_path).glob("*.xlsx")) + list(Path(folder_path).glob("*.xls"))
        
        if not excel_files:
            print(f"В папке {folder_path} не найдено Excel файлов")
            return False
        
        self.file_list = excel_files
        self.total_files = len(excel_files)
        
        print(f"\nНайдено {len(excel_files)} файлов:")
        for f in self.file_list:
            print(f"  - {f.name}")
        
        for file_path in self.file_list:
            print(f"\nОбработка: {file_path.name}")
            
            # Извлекаем шифр документа из имени файла
            doc_code_match = re.search(r'(\d+\.\d+-[A-Z0-9]+-[A-Z0-9]+-[A-Z0-9]+-[A-Z0-9]+\.[A-Z0-9]+)', file_path.stem)
            doc_code = doc_code_match.group(1) if doc_code_match else file_path.stem
            
            comments = self.parse_comment_review_sheet(file_path)
            
            if comments:
                print(f"  Найдено замечаний: {len(comments)}")
                self.total_remarks += len(comments)
                
                doc_data = self.remarks_data[doc_code]
                doc_data['total'] = len(comments)
                
                for comment in comments:
                    doc_data['by_status'][comment['status']] += 1
                    doc_data['details'].append(comment)
            else:
                print(f"  Замечаний не найдено")
                # Даже если нет замечаний, создаём запись о документе
                self.remarks_data[doc_code]['total'] = 0
        
        return True
    
    def create_summary_table(self, output_file="сводка_замечаний.xlsx"):
        """Создание сводной таблицы"""
        wb = Workbook()
        
        # Стили
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        total_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        warning_fill = PatternFill(start_color="FFE6E6", end_color="FFE6E6", fill_type="solid")
        border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        wrap_alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
        
        # ========== ЛИСТ 1: Сводка по документам ==========
        ws_summary = wb.active
        ws_summary.title = "Сводка по документам"
        
        headers = [
            '№ п/п', 'Шифр документа', 'Всего замечаний', 
            'Open', 'Closed', 'Процент устранения'
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws_summary.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        row = 2
        for idx, (doc_code, doc_data) in enumerate(sorted(self.remarks_data.items()), 1):
            open_count = doc_data['by_status'].get('Open', 0)
            closed_count = doc_data['by_status'].get('Closed', 0)
            total = doc_data['total']
            percent_closed = (closed_count / total * 100) if total > 0 else 100
            
            ws_summary.cell(row=row, column=1, value=idx).border = border
            ws_summary.cell(row=row, column=2, value=doc_code).border = border
            ws_summary.cell(row=row, column=3, value=total).border = border
            ws_summary.cell(row=row, column=4, value=open_count).border = border
            ws_summary.cell(row=row, column=5, value=closed_count).border = border
            percent_cell = ws_summary.cell(row=row, column=6, value=f"{percent_closed:.1f}%")
            percent_cell.border = border
            
            # Подсветка строк с открытыми замечаниями
            if open_count > 0:
                for col in range(1, 7):
                    ws_summary.cell(row=row, column=col).fill = warning_fill
            
            row += 1
        
        # Итоговая строка
        total_row = row
        for col in range(1, 7):
            cell = ws_summary.cell(row=total_row, column=col)
            cell.border = border
            cell.fill = total_fill
            cell.font = Font(bold=True)
        
        ws_summary.cell(row=total_row, column=1, value="ИТОГО:")
        ws_summary.cell(row=total_row, column=3, value=f"=SUM(C2:C{total_row-1})")
        ws_summary.cell(row=total_row, column=4, value=f"=SUM(D2:D{total_row-1})")
        ws_summary.cell(row=total_row, column=5, value=f"=SUM(E2:E{total_row-1})")
        
        # ========== ЛИСТ 2: Все замечания с текстом ==========
        ws_details = wb.create_sheet("Все замечания")
        
        detail_headers = [
            '№ п/п', 'Шифр документа', '№ замечания', 
            'Тип замечания', 'Текст замечания', 'Статус', 'Принято'
        ]
        
        for col, header in enumerate(detail_headers, 1):
            cell = ws_details.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        row = 2
        remark_idx = 1
        for doc_code, doc_data in sorted(self.remarks_data.items()):
            for comment in doc_data['details']:
                ws_details.cell(row=row, column=1, value=remark_idx).border = border
                ws_details.cell(row=row, column=2, value=doc_code).border = border
                ws_details.cell(row=row, column=3, value=comment['comment_no']).border = border
                ws_details.cell(row=row, column=4, value=comment['type']).border = border
                
                # Ячейка с текстом замечания
                cell_comment = ws_details.cell(row=row, column=5, value=comment['text'])
                cell_comment.border = border
                cell_comment.alignment = wrap_alignment
                
                ws_details.cell(row=row, column=6, value=comment['status']).border = border
                ws_details.cell(row=row, column=7, value=comment['acceptance']).border = border
                
                # Автовысота строки
                comment_lines = comment['text'].count('\n') + 1
                ws_details.row_dimensions[row].height = min(60, max(15, comment_lines * 12))
                row += 1
                remark_idx += 1
        
        # ========== ЛИСТ 3: Статистика по типам замечаний ==========
        ws_types = wb.create_sheet("Типы замечаний")
        
        type_stats = defaultdict(int)
        for doc_data in self.remarks_data.values():
            for comment in doc_data['details']:
                type_stats[comment['type']] += 1
        
        type_headers = ['№ п/п', 'Тип замечания', 'Количество', 'Процент']
        for col, header in enumerate(type_headers, 1):
            cell = ws_types.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
        
        row = 2
        for idx, (remark_type, count) in enumerate(sorted(type_stats.items(), key=lambda x: x[1], reverse=True), 1):
            percentage = (count / self.total_remarks * 100) if self.total_remarks > 0 else 0
            ws_types.cell(row=row, column=1, value=idx).border = border
            ws_types.cell(row=row, column=2, value=remark_type).border = border
            ws_types.cell(row=row, column=3, value=count).border = border
            ws_types.cell(row=row, column=4, value=f"{percentage:.1f}%").border = border
            row += 1
        
        # Итог для типов замечаний
        if type_stats:
            total_row_types = row
            ws_types.cell(row=total_row_types, column=1, value="ИТОГО:").border = border
            ws_types.cell(row=total_row_types, column=3, value=f"=SUM(C2:C{total_row_types-1})").border = border
            ws_types.cell(row=total_row_types, column=1).fill = total_fill
            ws_types.cell(row=total_row_types, column=3).fill = total_fill
            ws_types.cell(row=total_row_types, column=1).font = Font(bold=True)
            ws_types.cell(row=total_row_types, column=3).font = Font(bold=True)
        
        # ========== ЛИСТ 4: Статус замечаний ==========
        ws_status = wb.create_sheet("Статус замечаний")
        
        status_headers = ['Статус', 'Количество', 'Процент']
        for col, header in enumerate(status_headers, 1):
            cell = ws_status.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
        
        total_open = sum(doc['by_status'].get('Open', 0) for doc in self.remarks_data.values())
        total_closed = sum(doc['by_status'].get('Closed', 0) for doc in self.remarks_data.values())
        total_unknown = sum(doc['by_status'].get('Unknown', 0) for doc in self.remarks_data.values())
        
        status_data = [
            ('Open (Открыто)', total_open),
            ('Closed (Закрыто)', total_closed),
            ('Unknown (Не определен)', total_unknown)
        ]
        
        row = 2
        for status_name, count in status_data:
            percentage = (count / self.total_remarks * 100) if self.total_remarks > 0 else 0
            ws_status.cell(row=row, column=1, value=status_name).border = border
            ws_status.cell(row=row, column=2, value=count).border = border
            ws_status.cell(row=row, column=3, value=f"{percentage:.1f}%").border = border
            row += 1
        
        # Настройка ширины столбцов
        for ws in [ws_summary, ws_details, ws_types, ws_status]:
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Для листа с текстами делаем шире столбец с текстом
        ws_details.column_dimensions['E'].width = 60
        
        wb.save(output_file)
        print(f"\n✅ Сводная таблица сохранена: {output_file}")
        return output_file
    
    def print_statistics(self):
        """Вывод статистики в консоль"""
        print("\n" + "="*70)
        print("📊 СТАТИСТИКА ЗАМЕЧАНИЙ ПО КОМПЛЕКТУ ДОКУМЕНТАЦИИ")
        print("="*70)
        
        print(f"\n📁 Обработано файлов: {self.total_files}")
        print(f"📝 Всего замечаний: {self.total_remarks}")
        print(f"📄 Документов с замечаниями: {len([d for d in self.remarks_data.values() if d['total'] > 0])}")
        print(f"📑 Документов без замечаний: {len([d for d in self.remarks_data.values() if d['total'] == 0])}")
        
        # Топ документов
        if self.total_remarks > 0:
            print("\n📋 ТОП-5 документов с наибольшим количеством замечаний:")
            sorted_docs = sorted(self.remarks_data.items(), key=lambda x: x[1]['total'], reverse=True)[:5]
            for doc_code, doc_data in sorted_docs:
                if doc_data['total'] > 0:
                    print(f"  • {doc_code[:60]}... - {doc_data['total']} замечаний")
        
        # Статус замечаний
        total_open = sum(doc['by_status'].get('Open', 0) for doc in self.remarks_data.values())
        total_closed = sum(doc['by_status'].get('Closed', 0) for doc in self.remarks_data.values())
        
        print(f"\n✅ Статус устранения замечаний:")
        if self.total_remarks > 0:
            print(f"  • Открыто (Open): {total_open} ({total_open/self.total_remarks*100:.1f}%)")
            print(f"  • Закрыто (Closed): {total_closed} ({total_closed/self.total_remarks*100:.1f}%)")
            print(f"  • Процент устранения: {total_closed/self.total_remarks*100:.1f}%")
        else:
            print(f"  • Замечаний нет")
        
        # Типы замечаний
        type_stats = defaultdict(int)
        for doc_data in self.remarks_data.values():
            for comment in doc_data['details']:
                type_stats[comment['type']] += 1
        
        if type_stats:
            print(f"\n🏷️ Распределение по типам замечаний:")
            for remark_type, count in sorted(type_stats.items(), key=lambda x: x[1], reverse=True):
                percentage = (count / self.total_remarks * 100) if self.total_remarks > 0 else 0
                bar = "█" * int(percentage / 2)
                print(f"  {remark_type:<25}: {count:>3} ({percentage:>5.1f}%) {bar}")

def main():
    """Основная функция"""
    print("\n" + "="*70)
    print("🔍 ПРОГРАММА ПОДСЧЁТА ЗАМЕЧАНИЙ К КОМПЛЕКТУ РАБОЧЕЙ ДОКУМЕНТАЦИИ")
    print("="*70)
    
    folder_path = input("\n📂 Введите путь к папке с Excel файлами: ").strip()
    
    if not folder_path:
        folder_path = os.getcwd()
        print(f"Используется текущая папка: {folder_path}")
    
    if not os.path.exists(folder_path):
        print(f"❌ Ошибка: Папка '{folder_path}' не существует!")
        input("\nНажмите Enter для выхода...")
        return
    
    counter = DocumentRemarksCounter()
    
    if counter.process_folder(folder_path):
        counter.print_statistics()
        
        output_file = input("\n💾 Введите имя выходного файла (Enter для 'сводка_замечаний.xlsx'): ").strip()
        if not output_file:
            output_file = "сводка_замечаний.xlsx"
        
        # Добавляем расширение .xlsx если его нет
        if not output_file.endswith('.xlsx'):
            output_file += '.xlsx'
        
        counter.create_summary_table(output_file)
        print("\n✨ Готово! Сводная таблица создана успешно.")
    else:
        print("❌ Не удалось обработать файлы.")
    
    input("\nНажмите Enter для выхода...")

if __name__ == "__main__":
    main()